#!/usr/bin/env python3
"""Verify acceptance docx: TOC clean, forms filled, AI sections in chapter bodies."""
from __future__ import annotations

import re
import sys
from pathlib import Path

BACKEND = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND))

from docx import Document

from app.services.word import is_heading_style_name, is_toc_entry_text, paragraph_is_toc_like

MIN_AI_PARA_CHARS = 80
MIN_SECTION_BODY_CHARS = 60
STUB_PATTERNS = ("待填写", "TBD", "xxx")
AI_TITLES = [
    "竞标函",
    "法定代表人身份证明",
    "资格证明资料",
    "竞标人基本情况表",
    "履约信用证明",
    "其他材料",
    "竞标物资技术规格书",
    "生产组织供应能力分析表",
    "拟投入本项目的主要生产设备、检验设备表",
    "关于产品符合相关标准的承诺书",
]
AI_BODY_MARKERS = (
    "针对招标项目「",
    "华远智通",
)
EXPECTED_FORM_SNIPPETS = (
    ("华远智通科技有限公司", "bidder_name"),
    ("XCZHJT-2026-1027-A", "tender_no"),
    ("王明", "legal_name"),
    ("9800000", "bid_amount_lower"),
)
LEGACY_TEMPLATE_MARKERS = (
    "和远智能科技股份有限公司",
    "CRCEBG5-2026-CQSNTP011",
)


def norm_len(text: str) -> int:
    return len(re.sub(r"\s+", "", text or ""))


def _para_toc_like(style_name: str, text: str) -> bool:
    return paragraph_is_toc_like(style_name, text)


def _is_ai_substantive_text(text: str) -> bool:
    if any(m in text for m in AI_BODY_MARKERS):
        return True
    return text.startswith("关于「") and ("作如下说明" in text or "内容：" in text)


def _collect_text(doc: Document) -> str:
    parts: list[str] = []
    for p in doc.paragraphs:
        if p.text:
            parts.append(p.text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    parts.append(cell.text)
    return "\n".join(parts)


def _section_has_body(doc: Document, title: str) -> tuple[bool, int]:
    """标题段落后有实质性正文（LLM 撰写不必含「针对招标项目」套话）。"""
    paras = doc.paragraphs
    for i, p in enumerate(paras):
        st = p.style.name if p.style else ""
        text = (p.text or "").strip()
        if title not in text:
            continue
        if not is_heading_style_name(st) and title != text:
            continue
        for j in range(i + 1, min(i + 6, len(paras))):
            nxt = paras[j]
            nst = nxt.style.name if nxt.style else ""
            body = (nxt.text or "").strip()
            if not body or _para_toc_like(nst, body):
                continue
            if is_heading_style_name(nst) and title not in body:
                break
            n = norm_len(body)
            if n >= MIN_SECTION_BODY_CHARS and not any(s in body for s in STUB_PATTERNS):
                return True, n
    hits = [
        t
        for t in (
            p.text.strip()
            for p in doc.paragraphs
            if p.text and not _para_toc_like(p.style.name if p.style else "", p.text.strip())
        )
        if title in t
        and norm_len(t) >= MIN_AI_PARA_CHARS
        and _is_ai_substantive_text(t)
    ]
    if hits:
        return True, norm_len(hits[0])
    return False, 0


def main() -> int:
    path = Path(sys.argv[1] if len(sys.argv) > 1 else "/cursor/stores/self/media/acceptance-full-bid-sample.docx")
    doc = Document(str(path))
    full = _collect_text(doc)
    total = norm_len(full)
    ai_paras = [
        p.text.strip()
        for p in doc.paragraphs
        if p.text
        and _is_ai_substantive_text(p.text)
        and not _para_toc_like(p.style.name if p.style else "", p.text.strip())
    ]
    toc_ai_pollution = [
        p.text.strip()[:80]
        for p in doc.paragraphs
        if p.text
        and _para_toc_like(p.style.name if p.style else "", p.text.strip())
        and _is_ai_substantive_text(p.text)
    ]
    toc_substantive = [
        (p.style.name if p.style else "", p.text.strip()[:80])
        for p in doc.paragraphs
        if p.text
        and _para_toc_like(p.style.name if p.style else "", p.text.strip())
        and norm_len(p.text) > 45
        and not is_toc_entry_text(p.text)
    ]
    form_hits = {label: (snippet in full) for snippet, label in EXPECTED_FORM_SNIPPETS}
    legacy_left = [m for m in LEGACY_TEMPLATE_MARKERS if m in full]

    failures: list[str] = []
    report_lines = [
        f"DOCX: {path}",
        f"paragraphs_total={len(doc.paragraphs)}",
        f"total_chars={total}",
        f"ai_substantive_paragraphs={len(ai_paras)}",
        f"toc_ai_pollution={len(toc_ai_pollution)}",
        f"toc_substantive_paragraphs={len(toc_substantive)}",
        "",
        "form_fields:",
    ]
    for snippet, label in EXPECTED_FORM_SNIPPETS:
        flag = "OK" if form_hits[label] else "FAIL"
        report_lines.append(f"  {flag} {label} ({snippet[:24]}...)")
        if not form_hits[label]:
            failures.append(f"form:{label}")

    report_lines.append("")
    for title in AI_TITLES:
        ok, sample_len = _section_has_body(doc, title)
        flag = "OK" if ok else "FAIL"
        report_lines.append(f"{flag} {title} (body_chars={sample_len})")
        if not ok:
            failures.append(title)

    if toc_ai_pollution:
        failures.append("toc_ai_pollution")
    if toc_substantive:
        failures.append("toc_substantive_body")
    if legacy_left:
        failures.append(f"legacy_template_text:{legacy_left}")

    print("\n".join(report_lines))
    if failures or total < 5000:
        if total < 5000:
            failures.append(f"total_chars<{5000}")
        print("\nFAILED:", failures, file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
